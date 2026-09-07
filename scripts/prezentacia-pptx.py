#!/usr/bin/env python3
"""Zloží PPTX prezentáciu VESMA zo šablóny INOVIA a textov v docs/prezentacia-snimky.md.

Jednorazový nástroj pre prezentácie v subregiónoch (docs/prezentacia-scenar.md,
kap. 7). Nie je súčasťou aplikácie a nespúšťa sa v CI.

Šablóna INOVIA nie je postavená na placeholderoch — každá snímka je ručne
zložená z pozadia, farebných oválov, log a textových polí. Preto sa nové
snímky **klonujú zo vzorových snímok** šablóny a mení sa v nich len text.
Nová snímka tak má presne to isté pozadie, logá, fonty aj farby ako originál.

    pip install python-pptx
    python3 scripts/prezentacia-pptx.py --sablona SABLONA.pptx --iba-zoznam
    python3 scripts/prezentacia-pptx.py --sablona SABLONA.pptx \
        --vystup VESMA_prezentacia.pptx --uvod 1,2,3 --zaver 17

`--uvod` a `--zaver` sú čísla snímok šablóny, ktoré sa nechajú nezmenené na
začiatku a na konci. Ostatné snímky šablóny sa do výstupu nedostanú — slúžia
len ako vzory.
"""
from __future__ import annotations

import argparse
import copy
import json
import re
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

KLUCE_TEXT = ("id", "rozlozenie", "nadpis", "podnadpis", "vizual", "qr", "poznamky")

# Vzorové snímky šablóny INOVIA_Samosprávy (Prezentacia_Žarnay_samospráva.pptx).
# `snimka` je číslo snímky v šablóne od 1, `pole` názov textového poľa v nej,
# `nadpis`/`podnadpis` indexy odsekov s nadpisom, `zoznam_od` index prvého
# odseku zoznamu a `vzor_riadku` odsek, ktorého formátovanie sa kopíruje pri
# pridávaní ďalších riadkov.
# `velkost_riadku` je veľkosť písma riadkov zoznamu v bodoch — vzorové snímky
# majú kratšie texty než naše, takže bez zmenšenia by text pretiekol zo snímky.
VZORY_INOVIA = {
    "titulna":       {"snimka": 4,  "pole": "BlokTextu 7", "nadpis": 0, "podnadpis": 1, "zoznam_od": 3, "vzor_riadku": 3, "velkost_riadku": 22},
    "obsah":         {"snimka": 7,  "pole": "BlokTextu 7", "nadpis": 0, "zoznam_od": 2, "vzor_riadku": 2, "velkost_riadku": 20},
    "obsah_obrazok": {"snimka": 7,  "pole": "BlokTextu 7", "nadpis": 0, "zoznam_od": 2, "vzor_riadku": 2, "velkost_riadku": 18, "obrazok": True},
    "zaver":         {"snimka": 12, "pole": "BlokTextu 7", "nadpis": 0, "zoznam_od": 3, "vzor_riadku": 3, "velkost_riadku": 24},
}


# ---------------------------------------------------------------- obsah snímok

def nacitaj_snimky(cesta: Path) -> list[dict]:
    """Rozparsuje docs/prezentacia-snimky.md na zoznam snímok."""
    bloky = cesta.read_text(encoding="utf-8").split("\n---\n")[1:]
    snimky: list[dict] = []
    for blok in bloky:
        s: dict = {"odrazky": []}
        kluc = None
        for riadok in blok.splitlines():
            if riadok.startswith("- "):
                s["odrazky"].append(bez_uvodzoviek(riadok[2:].strip()))
                continue
            m = re.match(r"^([a-z_]+):\s*(.*)$", riadok)
            if m and m.group(1) in KLUCE_TEXT + ("odrazky",):
                kluc = m.group(1)
                if kluc != "odrazky":
                    s[kluc] = bez_uvodzoviek(m.group(2).strip())
                continue
            if kluc in KLUCE_TEXT and riadok.strip():
                s[kluc] = (s.get(kluc, "") + " " + riadok.strip()).strip()
        if s.get("id"):
            snimky.append(s)
    return snimky


def bez_uvodzoviek(hodnota: str) -> str:
    if len(hodnota) >= 2 and hodnota[0] == hodnota[-1] == '"':
        return hodnota[1:-1]
    return hodnota


# ------------------------------------------------------------ práca so snímkami

def klonuj_snimku(prez, zdroj):
    """Vytvorí novú snímku ako presnú kópiu zdrojovej (aj s obrázkami)."""
    ciel = prez.slides.add_slide(zdroj.slide_layout)
    for tvar in list(ciel.shapes):
        tvar._element.getparent().remove(tvar._element)

    # Vzťahy na obrázky a odkazy — rId sa v novej snímke môže líšiť, preto mapa.
    mapa: dict[str, str] = {}
    for rid, vztah in zdroj.part.rels.items():
        if vztah.reltype == RT.SLIDE_LAYOUT:
            continue
        if vztah.is_external:
            mapa[rid] = ciel.part.rels.get_or_add_ext_rel(vztah.reltype, vztah.target_ref)
        else:
            mapa[rid] = ciel.part.rels.get_or_add(vztah.reltype, vztah.target_part)

    # Pozadie snímky, ak ho zdroj definuje mimo tvarov.
    pozadie = zdroj._element.cSld.find(qn("p:bg"))
    if pozadie is not None:
        ciel._element.cSld.insert(0, copy.deepcopy(pozadie))

    strom = ciel.shapes._spTree
    for tvar in zdroj.shapes:
        prvok = copy.deepcopy(tvar._element)
        for uzol in prvok.iter():
            for atribut in (qn("r:embed"), qn("r:link"), qn("r:id")):
                stara = uzol.get(atribut)
                if stara in mapa:
                    uzol.set(atribut, mapa[stara])
        strom.append(prvok)
    return ciel


def najdi_pole(snimka, nazov: str):
    """Textové pole podľa názvu; keď nie je, tak najväčšie textové pole."""
    for tvar in snimka.shapes:
        if tvar.name == nazov and tvar.has_text_frame:
            return tvar
    kandidati = [t for t in snimka.shapes if t.has_text_frame and t.text_frame.text.strip()]
    return max(kandidati, key=lambda t: t.width * t.height) if kandidati else None


def nastav_odsek(odsek, text: str) -> None:
    """Prepíše text odseku a nechá formátovanie prvého behu."""
    if odsek.runs:
        odsek.runs[0].text = text
        for beh in odsek.runs[1:]:
            beh._r.getparent().remove(beh._r)
    else:
        odsek.text = text


def napln_pole(tvar, nadpis: str, podnadpis: str, riadky: list[str], vzor: dict) -> None:
    """Prepíše text vo vzorovom poli: nadpis, podnadpis a zoznam riadkov."""
    telo = tvar.text_frame._txBody
    odseky = tvar.text_frame.paragraphs

    if nadpis and vzor.get("nadpis") is not None and vzor["nadpis"] < len(odseky):
        nastav_odsek(odseky[vzor["nadpis"]], nadpis)
    if podnadpis and vzor.get("podnadpis") is not None and vzor["podnadpis"] < len(odseky):
        nastav_odsek(odseky[vzor["podnadpis"]], podnadpis)

    zaciatok = vzor.get("zoznam_od")
    if zaciatok is None:
        return

    vzorovy_index = min(vzor.get("vzor_riadku", zaciatok), len(odseky) - 1)
    vzorovy_prvok = copy.deepcopy(odseky[vzorovy_index]._p)

    # Odseky od `zaciatok` nižšie zahodíme a nahradíme našimi riadkami.
    for odsek in odseky[zaciatok:]:
        telo.remove(odsek._p)
    velkost = vzor.get("velkost_riadku")
    for riadok in riadky:
        telo.append(copy.deepcopy(vzorovy_prvok))
        odsek = tvar.text_frame.paragraphs[-1]
        nastav_odsek(odsek, riadok)
        if velkost:
            for beh in odsek.runs:
                beh.font.size = Pt(velkost)


def pridaj_miesto_na_obrazok(snimka, tvar_textu, popis: str, sirka: int, vyska: int) -> None:
    """Zúži textové pole a napravo doplní rámček na obrázok."""
    tvar_textu.width = Emu(int(sirka * 0.46))

    lavo, hore = int(sirka * 0.52), int(vyska * 0.26)
    ramcek = snimka.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Emu(lavo), Emu(hore), Emu(int(sirka * 0.40)), Emu(int(vyska * 0.52)),
    )
    ramcek.name = "MIESTO NA OBRÁZOK"
    ramcek.fill.solid()
    ramcek.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    ramcek.fill.transparency = 0.15
    ramcek.line.color.rgb = RGBColor(0xF3, 0xCE, 0x3C)
    ramcek.line.width = Pt(1.5)
    ramec = ramcek.text_frame
    ramec.word_wrap = True
    odsek = ramec.paragraphs[0]
    odsek.text = f"MIESTO NA OBRÁZOK\n{popis}"
    for beh in odsek.runs:
        beh.font.size = Pt(12)
        beh.font.name = "Nunito"
        beh.font.color.rgb = RGBColor(0x33, 0x33, 0x33)


def vyrob_qr(url: str, cesta: Path) -> Path:
    """Vygeneruje QR kód s bielym podkladom, aby bol čitateľný na farbe."""
    import qrcode

    obrazok = qrcode.make(url, box_size=12, border=3)
    obrazok.save(cesta)
    return cesta


def pridaj_qr(snimka, cesta: Path, lavo: int, hore: int, velkost: int) -> None:
    snimka.shapes.add_picture(str(cesta), Emu(lavo), Emu(hore), Emu(velkost), Emu(velkost))


def presun_snimku(prez, z_indexu: int, na_index: int) -> None:
    zoznam = prez.slides._sldIdLst
    prvok = list(zoznam)[z_indexu]
    zoznam.remove(prvok)
    zoznam.insert(na_index, prvok)


def zmaz_snimky(prez, indexy: set[int]) -> None:
    """Zmaže snímky podľa indexov od 0."""
    zoznam = prez.slides._sldIdLst
    prvky = list(zoznam)
    for i in sorted(indexy, reverse=True):
        prez.part.drop_rel(prvky[i].rId)
        zoznam.remove(prvky[i])


# ------------------------------------------------------------------- zostavenie

def vypis_sablonu(prez) -> None:
    print(f"Snímok v šablóne: {len(prez.slides._sldIdLst)}")
    for i, snimka in enumerate(prez.slides, start=1):
        texty = [
            t.text_frame.text.strip().replace("\n", " | ")[:70]
            for t in snimka.shapes
            if t.has_text_frame and t.text_frame.text.strip()
        ]
        print(f"  {i:2d}: {' / '.join(texty)[:110] or '(bez textu)'}")


def poskladaj(args) -> None:
    prez = Presentation(args.sablona)
    if args.iba_zoznam:
        vypis_sablonu(prez)
        return

    vzory = dict(VZORY_INOVIA)
    if args.vzory:
        for typ, hodnoty in json.loads(args.vzory).items():
            vzory.setdefault(typ, {}).update(hodnoty)

    snimky = nacitaj_snimky(Path(args.obsah))
    if not snimky:
        sys.exit(f"V {args.obsah} som nenašiel ani jednu snímku.")

    pocet_vzorov = len(prez.slides._sldIdLst)
    uvod = [int(c) for c in args.uvod.split(",") if c.strip()]
    zaver = [int(c) for c in args.zaver.split(",") if c.strip()]
    sirka, vyska = prez.slide_width, prez.slide_height

    print(f"Šablóna: {pocet_vzorov} snímok · úvodné {uvod or '—'} · záverečné {zaver or '—'}")
    for s in snimky:
        vzor = vzory.get(s["rozlozenie"])
        if vzor is None:
            sys.exit(f"{s['id']}: neznáme rozloženie {s['rozlozenie']!r}.")
        zdroj = prez.slides[vzor["snimka"] - 1]
        nova = klonuj_snimku(prez, zdroj)

        pole = najdi_pole(nova, vzor.get("pole", ""))
        if pole is None:
            sys.exit(f"{s['id']}: vo vzorovej snímke {vzor['snimka']} nie je textové pole.")
        napln_pole(pole, s.get("nadpis", ""), s.get("podnadpis", ""), s.get("odrazky", []), vzor)

        poznamky = s.get("poznamky", "")
        if s.get("qr"):
            cesta = Path(args.vystup).with_name(f"qr-{s['id']}.png")
            vyrob_qr(s["qr"], cesta)
            if vzor.get("obrazok"):
                pole.width = Emu(int(sirka * 0.46))
                pridaj_qr(nova, cesta, int(sirka * 0.60), int(vyska * 0.28), int(vyska * 0.44))
            else:
                pridaj_qr(nova, cesta, int(sirka * 0.04), int(vyska * 0.755), int(vyska * 0.20))
            poznamky = f"QR kód vedie na {s['qr']}.\n\n{poznamky}".strip()
        elif s.get("vizual"):
            if vzor.get("obrazok"):
                pridaj_miesto_na_obrazok(nova, pole, s["vizual"], sirka, vyska)
            else:
                poznamky = f"OBRÁZOK: {s['vizual']}\n\n{poznamky}".strip()
        if poznamky:
            nova.notes_slide.notes_text_frame.text = poznamky

        print(f"  {s['id']:3s} ← vzor {vzor['snimka']:2d} · {s.get('nadpis', '')[:52]}")

    # Poradie: úvodné snímky šablóny, naše snímky, záverečné snímky šablóny.
    for cielovy, cislo in enumerate(uvod):
        presun_snimku(prez, cislo - 1, cielovy)
    for cislo in zaver:
        presun_snimku(prez, cislo - 1, len(prez.slides._sldIdLst) - 1)

    ponechane = {c - 1 for c in uvod} | {c - 1 for c in zaver}
    zmaz_snimky(prez, set(range(pocet_vzorov)) - ponechane)

    # Záverečné snímky šablóny patria pred záložné snímky (id začínajúce Z).
    pocet_zalozných = sum(1 for s in snimky if s["id"].startswith("Z"))
    for _ in zaver:
        presun_snimku(prez, len(prez.slides._sldIdLst) - 1, len(prez.slides._sldIdLst) - 1 - pocet_zalozných)

    prez.save(args.vystup)
    print(f"\nHotovo: {args.vystup} ({len(prez.slides._sldIdLst)} snímok)")


def main() -> None:
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--sablona", required=True, help="PPTX šablóna INOVIA")
    p.add_argument("--vystup", default="VESMA_prezentacia.pptx")
    p.add_argument("--obsah", default="docs/prezentacia-snimky.md")
    p.add_argument("--uvod", default="", help="čísla snímok šablóny na úvod, napr. 1,2,3")
    p.add_argument("--zaver", default="", help="čísla snímok šablóny na záver, napr. 17")
    p.add_argument("--vzory", default="", help='úprava vzorov, napr. \'{"obsah":{"snimka":9}}\'')
    p.add_argument("--iba-zoznam", action="store_true", help="vypíše snímky šablóny s ich textom")
    poskladaj(p.parse_args())


if __name__ == "__main__":
    main()
